import unittest
import os
import openpyxl
from datetime import datetime
import tempfile
import json
import sys

# Garante que app/api/Services está no PYTHONPATH
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '../../app/api/Services')))

try:
    from preditiva_parser import process_preditiva_excel
except ImportError:
    process_preditiva_excel = None


class TestPreditivaParser(unittest.TestCase):
    def setUp(self):
        self.clima_file = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../CLIMA_20261610.xlsx'))
        self.template_file = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../Preditiva_parse.xlsx'))
        self.temp_dir = tempfile.TemporaryDirectory()
        self.output_file = os.path.join(self.temp_dir.name, 'output_test.xlsx')

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_parser_exists(self):
        self.assertIsNotNone(process_preditiva_excel, "Função process_preditiva_excel deve existir")

    def test_process_preditiva_matching(self):
        stats = process_preditiva_excel(self.clima_file, self.template_file, self.output_file)
        
        self.assertTrue(os.path.exists(self.output_file), "Arquivo de saída deve ser gerado")
        self.assertIsInstance(stats, dict)
        self.assertIn("total_rows", stats)
        self.assertIn("matched_rows", stats)
        self.assertIn("unmatched_rows", stats)
        self.assertEqual(stats["total_rows"], 304)
        self.assertGreaterEqual(stats["matched_rows"], 200)

        # Inspeciona o XLSX gerado
        wb = openpyxl.load_workbook(self.output_file)
        self.assertIn('PREVENTIVAS_PARSE', wb.sheetnames)
        ws = wb['PREVENTIVAS_PARSE']

        # Linha 2 (BMADTC / CLIMA - 01 / SPRINGER CARRIER - SILENTIA) -> dá match com linha 8 de CLIMA
        row2_status = ws.cell(2, 13).value
        self.assertEqual(row2_status, "FINALIZADA")
        
        row2_executor = ws.cell(2, 11).value
        self.assertEqual(row2_executor, "LEONARDO DE OLIVIERA VIANA")
        
        row2_conformidade = ws.cell(2, 12).value
        self.assertEqual(row2_conformidade, "OK")

        # Linha 3 (BMADTC / CLIMA - 02 / SPRINGER - MUNDIAL) -> dá match com linha 9 de CLIMA
        row3_status = ws.cell(3, 13).value
        self.assertEqual(row3_status, "FINALIZADA")

        row3_executor = ws.cell(3, 11).value
        self.assertEqual(row3_executor, "LEONARDO DE OLIVIERA VIANA")

        # Verifica preservação de abas adicionais
        self.assertIn('FEV', wb.sheetnames)
        self.assertIn('MAR', wb.sheetnames)

    def test_process_preditiva_with_unextended_temp_file(self):
        # Simula arquivo temporário do PHP sem extensão .xlsx
        import shutil
        unextended_clima = os.path.join(self.temp_dir.name, 'phpUploadTemp12345')
        shutil.copy(self.clima_file, unextended_clima)

        out = os.path.join(self.temp_dir.name, 'out_unextended.xlsx')
        stats = process_preditiva_excel(unextended_clima, self.template_file, out)
        self.assertTrue(stats["success"])
        self.assertEqual(stats["total_rows"], 304)


if __name__ == '__main__':
    unittest.main()
